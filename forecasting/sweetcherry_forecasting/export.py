"""Forecast exports in the v1 contract the admin upload accepts.

Workbook: sheets "Grade A", "Grade AA", "Grade AAA" with the columns
Date | Market | Low (P10) | Expected (P50) | High (P90) | Horizon Type | Horizon,
then an "About" sheet (last) whose first rows are machine-readable key/value
pairs. Dates are written as ISO text so no time zone can shift them.
JSON: the same rows per grade plus the same About values.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config as C
from .forecast import Forecast

GRADE_SHEETS = {"A": "Grade A", "AA": "Grade AA", "AAA": "Grade AAA"}
HEADER = ["Date", "Market", "Low (P10)", "Expected (P50)", "High (P90)", "Horizon Type", "Horizon"]
ABOUT_SHEET = "About"
NOTICE = ("SIMULATED DATA: these prices come from a simulation, not from real markets. The backtest figures check "
          "the code on simulated data; they are not real-world accuracy.")
METRIC_COLUMNS = [
    ("horizon", None),
    ("n", "n"),
    ("MAE (MAD/kg)", "mae"),
    ("MAPE %", "mape_pct"),
    ("P10-P90 coverage %", "coverage_pct"),
    ("mean P10-P90 width (MAD/kg)", "mean_width"),
    ("naive MAE (MAD/kg)", "naive_mae"),
    ("seasonal-naive MAE (MAD/kg)", "seasonal_naive_mae"),
]


def about_values(forecast: Forecast, *, generated_at: str, model_version: str, source_name: str, seed: int,
                 generator_version: str, input_workbook: str, first_evaluated: int,
                 test_seasons: tuple[int, ...]) -> list[tuple[str, str]]:
    cutoff = forecast.cutoff.date().isoformat()
    return [
        ("provenance", "synthetic"),
        ("dataCutoff", cutoff),
        ("generatedAt", generated_at),
        ("modelVersion", model_version),
        ("sourceName", source_name),
        ("seed", str(seed)),
        ("generatorVersion", generator_version),
        ("notice", NOTICE),
        ("snapshot", f"Fixed demonstration snapshot with data up to {cutoff}; not current market data."),
        ("units", "MAD/kg, wholesale"),
        ("horizons", "1w to 8w after the cutoff week (short_term)"),
        ("markets", "; ".join(f"{market.code} {market.name}" for market in C.MARKETS)),
        ("inputWorkbook", input_workbook),
        ("rowCounts", ", ".join(f"{grade} {len(forecast.rows.get(grade, []))}" for grade in C.GRADES)),
        ("calibration", f"P10-P90 widened per horizon to cover 80% of out-of-sample simulated weeks, seasons "
                        f"{first_evaluated} to {max(test_seasons)} (conformal)"),
        ("backtest", f"Rolling-season backtest on simulated seasons {min(test_seasons)} to {max(test_seasons)}; "
                     f"each season predicted by models trained only on earlier seasons"),
        ("forecastNote", forecast.note),
    ]


def _metric_rows(metrics: dict) -> list[list]:
    rows = []
    blocks = list(metrics["by_horizon"].items()) + [("all", metrics["all_horizons"])]
    for label, block in blocks:
        rows.append([label] + [round(block[key], 2) if key != "n" else block[key] for _, key in METRIC_COLUMNS[1:]])
    return rows


def write_forecast_workbook(forecast: Forecast, about: list[tuple[str, str]], metrics: dict, path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    book.properties.title = "SIMULATED cherry price forecast"
    book.properties.subject = "Forecast made from simulated data, not real market prices"

    for grade, title in GRADE_SHEETS.items():
        sheet = book.create_sheet(title)
        sheet.append(HEADER)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in forecast.rows.get(grade, []):
            sheet.append([row["date"], row["market"], row["low"], row["expected"], row["high"],
                          row["horizonType"], row["horizon"]])
        for column, width in enumerate([12, 9, 11, 15, 12, 13, 9], start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        for cells in sheet.iter_rows(min_row=2, min_col=3, max_col=5):
            for cell in cells:
                cell.number_format = "0.00"
        sheet.freeze_panes = "A2"

    sheet = book.create_sheet(ABOUT_SHEET)
    for key, value in about:
        sheet.append([key, value])
    for row in sheet.iter_rows(min_row=1, max_row=len(about), max_col=1):
        row[0].font = Font(bold=True)
    notice_row = next(index for index, (key, _) in enumerate(about, start=1) if key == "notice")
    for column in (1, 2):
        cell = sheet.cell(row=notice_row, column=column)
        cell.fill = PatternFill("solid", fgColor="FFF2CC")
        cell.font = Font(bold=True, color="7F3F00")
    sheet.cell(row=notice_row, column=2).alignment = Alignment(wrap_text=True)
    sheet.append([])
    sheet.append(["Backtest on SIMULATED data: a code check, not real-world accuracy"])
    sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
    sheet.append([name for name, _ in METRIC_COLUMNS])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    for row in _metric_rows(metrics):
        sheet.append(row)
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 110
    for column in range(3, len(METRIC_COLUMNS) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 18
    book.active = book.sheetnames.index(ABOUT_SHEET)  # opens on the provenance, stays last in order
    book.save(path)
    return path


def write_forecast_json(forecast: Forecast, about: list[tuple[str, str]], metrics: dict, path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "schema": "sweetcherry.forecastSheets.v1",
        "about": dict(about),
        "backtest": {
            "byHorizon": {label: _rounded(block) for label, block in metrics["by_horizon"].items()},
            "allHorizons": _rounded(metrics["all_horizons"]),
        },
        "grades": {grade: forecast.rows.get(grade, []) for grade in C.GRADES},
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def _rounded(block: dict) -> dict:
    return {key: (round(value, 3) if isinstance(value, float) else value) for key, value in block.items()}
