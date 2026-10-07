"""Excel workbook for the SIMULATED market data, and the reader used by the
pipeline. The workbook says on every sheet that the data are simulated."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import SOURCE_NAME
from . import config as C
from .synthetic import COLUMNS, SyntheticMarket

OBSERVATIONS_SHEET = "Observations"
BANNER_FILL = PatternFill("solid", fgColor="FFF2CC")
HEADER_FILL = PatternFill("solid", fgColor="E7ECEA")
BOLD = Font(bold=True)

DATA_DICTIONARY = [
    ("date", "date", "", "Trading day of the simulated bulletin (Monday to Saturday).", "2018-01-01 to the cutoff"),
    ("market_id", "text", "", "Wholesale market code.", "M01, M02, M03"),
    ("market_name", "text", "", "Wholesale market name.", "Casablanca, Rabat, Marrakech"),
    ("grade", "text", "", "Quality grade, AAA being the largest fruit.", "A, AA, AAA"),
    ("supply_type", "text", "", "Where the fruit on the line comes from.", "fresh_local, cold_storage, import"),
    ("main_origin", "text", "", "Region supplying most of the line's fruit; 'Import' for imported fruit.",
     ", ".join(region.name for region in C.REGIONS) + ", Import"),
    ("price_min_mad_kg", "number", "MAD/kg", "Lowest price on the line, rounded to 0.5 MAD.", "empty when closed"),
    ("price_modal_mad_kg", "number", "MAD/kg", "Most frequent (modal) price, rounded to 0.5 MAD. The forecasts use its weekly median.", "empty when closed"),
    ("price_max_mad_kg", "number", "MAD/kg", "Highest price on the line, rounded to 0.5 MAD.", "empty when closed"),
    ("arrivals_t", "number", "tonnes", "Simulated arrivals of that grade on that market and day.", "empty when closed"),
    ("status", "text", "", "Whether the bulletin was published or the market was closed.", "reported, closed_holiday"),
    ("quality_flag", "text", "", "Lines a checker would look at twice.", "empty, explained_outlier, suspect_entry"),
    ("remark", "text", "", "Reason for an outlier, the holiday name for a closure, or a data-entry warning.", "free text"),
]


def _banner(sheet, row: int, text: str, width: int) -> None:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    cell = sheet.cell(row=row, column=1, value=text)
    cell.font = Font(bold=True, color="7F3F00")
    cell.fill = BANNER_FILL
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[row].height = 34


def _header(sheet, row: int, values: list[str]) -> None:
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.font = BOLD
        cell.fill = HEADER_FILL


def _widths(sheet, widths: list[int]) -> None:
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def provenance(market: SyntheticMarket, generated_at: str) -> list[tuple[str, str]]:
    observations = market.observations
    reported = int((observations["status"] == "reported").sum())
    return [
        ("provenance", "synthetic"),
        ("sourceName", SOURCE_NAME),
        ("generatorVersion", market.generator_version),
        ("seed", str(market.seed)),
        ("periodStart", market.start.date().isoformat()),
        ("dataCutoff", market.cutoff.date().isoformat()),
        ("generatedAt", generated_at),
        ("units", "Prices in Moroccan dirhams per kilogram (MAD/kg, wholesale); arrivals in tonnes"),
        ("markets", "; ".join(f"{market_.code} {market_.name}" for market_ in C.MARKETS)),
        ("grades", ", ".join(C.GRADES)),
        ("rows", f"{len(observations)} lines: {reported} reported, {len(observations) - reported} holiday closures"),
        ("regenerate", f"python -m sweetcherry_forecasting generate --seed {market.seed} --start {market.start.date()} --cutoff {market.cutoff.date()}"),
    ]


def write_market_workbook(market: SyntheticMarket, path: Path, generated_at: str) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.properties.title = "SIMULATED wholesale cherry prices"
    book.properties.subject = "Simulated data, not real market prices"
    book.properties.creator = f"sweetcherry_forecasting generator {market.generator_version}"

    readme = book.active
    readme.title = "Read me"
    readme.cell(row=1, column=1, value="SweetCherry: simulated wholesale cherry prices (Casablanca, Rabat, Marrakech)").font = Font(bold=True, size=14)
    _banner(readme, 3, "SIMULATED DATA. A computer program generated these prices for testing and demonstration. "
            "They are not observations from any real market, and the events in 'Simulated events' did not happen.", 4)
    _header(readme, 5, ["Field", "Value"])
    for offset, (key, value) in enumerate(provenance(market, generated_at), start=6):
        readme.cell(row=offset, column=1, value=key)
        readme.cell(row=offset, column=2, value=value)
    notes_row = 6 + len(provenance(market, generated_at)) + 1
    for offset, text in enumerate([
        "Sheets: 'Data dictionary' explains each column; 'Assumptions' lists every simulation setting;",
        "'Simulated events' lists the invented shocks; 'Observations' holds the bulletin lines.",
        "Backtests on this data test the forecasting code. They say nothing about accuracy on a real market.",
    ]):
        readme.cell(row=notes_row + offset, column=1, value=text)
    _widths(readme, [20, 110, 12, 12])

    dictionary = book.create_sheet("Data dictionary")
    _banner(dictionary, 1, "Columns of the 'Observations' sheet. All values are SIMULATED.", 5)
    _header(dictionary, 2, ["column", "type", "unit", "description", "values"])
    for offset, row in enumerate(DATA_DICTIONARY, start=3):
        for column, value in enumerate(row, start=1):
            dictionary.cell(row=offset, column=column, value=value)
    _widths(dictionary, [20, 9, 9, 90, 48])

    assumptions = book.create_sheet("Assumptions")
    _banner(assumptions, 1, "Illustrative assumptions used by the simulator, not measured facts. Review them against real "
            "market data before relying on any of them.", 3)
    _header(assumptions, 2, ["topic", "assumption", "value"])
    for offset, row in enumerate(C.assumption_rows(), start=3):
        for column, value in enumerate(row, start=1):
            assumptions.cell(row=offset, column=column, value=value)
    _widths(assumptions, [18, 78, 92])

    events = book.create_sheet("Simulated events")
    _banner(events, 1, "Simulated events: shocks invented by the simulator. They did not happen and are not weather records.", 6)
    _header(events, 2, ["year", "event", "start", "end", "scope", "simulated effect"])
    for offset, record in enumerate(market.events.itertuples(index=False), start=3):
        values = [record.year, record.event, record.start.to_pydatetime(), record.end.to_pydatetime(), record.scope, record.simulated_effect]
        for column, value in enumerate(values, start=1):
            cell = events.cell(row=offset, column=column, value=value)
            if column in (3, 4):
                cell.number_format = "yyyy-mm-dd"
    _widths(events, [7, 16, 12, 12, 30, 90])

    sheet = book.create_sheet(OBSERVATIONS_SHEET)
    _banner(sheet, 1, f"SIMULATED DATA, not real market prices. Generator {market.generator_version}, seed {market.seed}, "
            f"cutoff {market.cutoff.date()}.", len(COLUMNS))
    _header(sheet, 2, COLUMNS)
    for offset, record in enumerate(market.observations.itertuples(index=False), start=3):
        for column, name in enumerate(COLUMNS, start=1):
            value = getattr(record, name)
            if name == "date":
                value = value.to_pydatetime()
            elif isinstance(value, float) and pd.isna(value):
                value = None
            elif value == "":
                value = None
            cell = sheet.cell(row=offset, column=column, value=value)
            if name == "date":
                cell.number_format = "yyyy-mm-dd"
            elif name.endswith("_mad_kg") or name == "arrivals_t":
                cell.number_format = "0.0"
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(COLUMNS))}{len(market.observations) + 2}"
    _widths(sheet, [11, 9, 12, 6, 13, 12, 15, 17, 15, 10, 15, 17, 58])

    book.active = 0
    book.save(path)
    return path


def read_observations(path: Path) -> pd.DataFrame:
    """Bulletin lines from a market workbook, typed like the generator output."""
    frame = pd.read_excel(path, sheet_name=OBSERVATIONS_SHEET, header=1)
    missing = [column for column in COLUMNS if column not in frame.columns]
    if missing:
        raise ValueError(f"{path}: missing columns {missing}")
    frame = frame[COLUMNS].copy()
    frame["date"] = pd.to_datetime(frame["date"]).dt.normalize()
    for column in ("market_id", "market_name", "grade", "supply_type", "main_origin", "status", "quality_flag", "remark"):
        frame[column] = frame[column].fillna("").astype(str)
    for column in ("price_min_mad_kg", "price_modal_mad_kg", "price_max_mad_kg", "arrivals_t"):
        frame[column] = pd.to_numeric(frame[column], errors="coerce")
    return frame


def read_provenance(path: Path) -> dict[str, str]:
    sheet = pd.read_excel(path, sheet_name="Read me", header=None)
    values = {}
    for key, value in zip(sheet.iloc[:, 0], sheet.iloc[:, 1]):
        if isinstance(key, str) and isinstance(value, (str, int, float)) and not pd.isna(value):
            values[key] = str(value)
    return values
