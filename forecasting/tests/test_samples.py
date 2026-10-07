"""The committed samples are reproducible from code and agree with each
other, with the backtest report and with docs/MODEL_CARDS.md."""

from __future__ import annotations

import json

import pandas as pd
from openpyxl import load_workbook

from sweetcherry_forecasting.cli import MODEL_CARDS, REPORTS, SAMPLES, command_run, forecast_name, market_workbook_name
from sweetcherry_forecasting import DEFAULT_CUTOFF, DEFAULT_SEED
from sweetcherry_forecasting.export import HEADER
from sweetcherry_forecasting.report import update_doc
from sweetcherry_forecasting.synthetic import generate
from sweetcherry_forecasting.workbook_raw import read_observations

MARKET_SAMPLE = SAMPLES / market_workbook_name(DEFAULT_SEED)
FORECAST_XLSX = SAMPLES / forecast_name(DEFAULT_CUTOFF, "xlsx")
FORECAST_JSON = SAMPLES / forecast_name(DEFAULT_CUTOFF, "json")
METRICS = REPORTS / "backtest_metrics.json"


def test_committed_market_sample_matches_the_generator():
    committed = read_observations(MARKET_SAMPLE)
    pd.testing.assert_frame_equal(committed, generate().observations, check_dtype=False)


def test_forecast_workbook_layout():
    book = load_workbook(FORECAST_XLSX)
    assert book.sheetnames == ["Grade A", "Grade AA", "Grade AAA", "About"]
    for name in book.sheetnames[:3]:
        assert [cell.value for cell in book[name][1]] == HEADER
    about = {row[0]: row[1] for row in book["About"].iter_rows(values_only=True) if row and row[0]}
    for key in ("provenance", "dataCutoff", "generatedAt", "modelVersion", "sourceName", "seed"):
        assert key in about, key
    assert about["provenance"] == "synthetic" and about["dataCutoff"] == DEFAULT_CUTOFF
    assert about["sourceName"] == "Synthetic market simulation"


def test_forecast_json_matches_the_workbook_and_report():
    payload = json.loads(FORECAST_JSON.read_text(encoding="utf-8"))
    book = load_workbook(FORECAST_XLSX)
    for grade in ("A", "AA", "AAA"):
        sheet_rows = [row for row in book[f"Grade {grade}"].iter_rows(min_row=2, values_only=True)]
        json_rows = [(r["date"], r["market"], r["low"], r["expected"], r["high"], r["horizonType"], r["horizon"])
                     for r in payload["grades"][grade]]
        assert sheet_rows == json_rows
    report = json.loads(METRICS.read_text(encoding="utf-8"))
    assert payload["backtest"]["allHorizons"] == report["metrics"]["all_horizons"]
    assert report["provenance"] == "synthetic"


def test_model_card_table_matches_the_report():
    assert update_doc(MODEL_CARDS, METRICS, check=True), "run: python -m sweetcherry_forecasting run"


def test_rerunning_the_pipeline_reproduces_the_committed_forecast(tmp_path):
    class Args:
        data = str(MARKET_SAMPLE)
        seed = DEFAULT_SEED
        cutoff = None
        samples = str(tmp_path)
        reports = str(tmp_path)
        no_docs = True

    command_run(Args())
    fresh = json.loads((tmp_path / FORECAST_JSON.name).read_text(encoding="utf-8"))
    committed = json.loads(FORECAST_JSON.read_text(encoding="utf-8"))
    for grade in ("A", "AA", "AAA"):
        new, old = pd.DataFrame(fresh["grades"][grade]), pd.DataFrame(committed["grades"][grade])
        assert new[["date", "market", "horizon"]].equals(old[["date", "market", "horizon"]])
        # LightGBM may differ in the last digits across platforms; prices must agree to a cent or two.
        assert (new[["low", "expected", "high"]] - old[["low", "expected", "high"]]).abs().max().max() <= 0.02
